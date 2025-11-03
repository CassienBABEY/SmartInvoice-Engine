"""Utilitaires pour la manipulation de fichiers DOCX."""
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches


def find_placeholders(docx_path: Path) -> list[str]:
    """
    Extrait tous les placeholders {{field}} d'un document DOCX.

    Note: Reconstitue le texte des runs pour gérer la fragmentation Word.
    Scanne paragraphes, tableaux, headers et footers.

    Args:
        docx_path: Chemin vers le fichier DOCX

    Returns:
        Liste unique des placeholders trouvés
    """
    from app.core.config import FIELD_PATTERN

    doc = Document(docx_path)
    placeholders = set()

    # Chercher dans les paragraphes
    for paragraph in doc.paragraphs:
        full_text = _get_paragraph_text(paragraph)
        matches = re.findall(FIELD_PATTERN, full_text)
        placeholders.update(matches)

    # Chercher dans les tableaux
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    full_text = _get_paragraph_text(paragraph)
                    matches = re.findall(FIELD_PATTERN, full_text)
                    placeholders.update(matches)

                # Chercher dans les tables imbriquées
                for nested_table in cell.tables:
                    placeholders.update(
                        _extract_from_table(nested_table, FIELD_PATTERN)
                    )

    # Chercher dans les headers
    for section in doc.sections:
        header = section.header
        for paragraph in header.paragraphs:
            full_text = _get_paragraph_text(paragraph)
            matches = re.findall(FIELD_PATTERN, full_text)
            placeholders.update(matches)

        for table in header.tables:
            placeholders.update(_extract_from_table(table, FIELD_PATTERN))

    # Chercher dans les footers
    for section in doc.sections:
        footer = section.footer
        for paragraph in footer.paragraphs:
            full_text = _get_paragraph_text(paragraph)
            matches = re.findall(FIELD_PATTERN, full_text)
            placeholders.update(matches)

        for table in footer.tables:
            placeholders.update(_extract_from_table(table, FIELD_PATTERN))

    # Chercher dans les text boxes via XML brut
    textbox_placeholders = _extract_from_textboxes(docx_path, FIELD_PATTERN)
    placeholders.update(textbox_placeholders)

    return sorted(list(placeholders))


def replace_placeholders(
    docx_path: Path, replacements: dict[str, str], output_path: Path
) -> None:
    """
    Remplace les placeholders dans un DOCX et sauvegarde le résultat.

    Remplace dans paragraphes, tableaux, headers et footers.
    Les placeholders non fournis ou vides sont supprimés du document.

    Args:
        docx_path: Chemin vers le template DOCX
        replacements: Dictionnaire {placeholder: valeur}
        output_path: Chemin de sortie pour le DOCX généré
    """
    doc = Document(docx_path)

    # Construire un dictionnaire complet avec tous les placeholders
    # Ceux non fournis ou vides seront remplacés par ""
    complete_replacements = {}
    all_placeholders = find_placeholders(docx_path)
    
    for placeholder in all_placeholders:
        # Si le placeholder est dans replacements et a une valeur
        if placeholder in replacements and replacements[placeholder]:
            complete_replacements[placeholder] = replacements[placeholder]
        else:
            # Sinon, le remplacer par une chaîne vide pour le supprimer
            complete_replacements[placeholder] = ""

    # Remplacer dans les paragraphes
    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, complete_replacements)

    # Remplacer dans les tableaux
    for table in doc.tables:
        _replace_in_table(table, complete_replacements)

    # Remplacer dans les headers
    for section in doc.sections:
        header = section.header
        for paragraph in header.paragraphs:
            _replace_in_paragraph(paragraph, complete_replacements)
        for table in header.tables:
            _replace_in_table(table, complete_replacements)

    # Remplacer dans les footers
    for section in doc.sections:
        footer = section.footer
        for paragraph in footer.paragraphs:
            _replace_in_paragraph(paragraph, complete_replacements)
        for table in footer.tables:
            _replace_in_table(table, complete_replacements)

    doc.save(output_path)

    # Remplacer dans les text boxes via XML
    _replace_in_textboxes_xml(output_path, complete_replacements)


def _extract_from_table(table, pattern: str) -> set[str]:
    """
    Extrait les placeholders d'un tableau.

    Args:
        table: Objet table python-docx
        pattern: Pattern regex à rechercher

    Returns:
        Set de placeholders trouvés
    """
    placeholders = set()
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                full_text = _get_paragraph_text(paragraph)
                matches = re.findall(pattern, full_text)
                placeholders.update(matches)
    return placeholders


def _extract_from_textboxes(docx_path: Path, pattern: str) -> set[str]:
    """
    Extrait les placeholders des text boxes via parsing XML.

    Args:
        docx_path: Chemin vers le fichier DOCX
        pattern: Pattern regex à rechercher

    Returns:
        Set de placeholders trouvés dans les text boxes
    """
    placeholders = set()

    try:
        with zipfile.ZipFile(docx_path, "r") as docx_zip:
            # Parser document.xml
            if "word/document.xml" in docx_zip.namelist():
                xml_content = docx_zip.read("word/document.xml").decode("utf-8")
                clean_text = _extract_text_from_xml(xml_content)
                matches = re.findall(pattern, clean_text)
                placeholders.update(matches)

            # Parser header*.xml
            for name in docx_zip.namelist():
                if name.startswith("word/header") and name.endswith(".xml"):
                    xml_content = docx_zip.read(name).decode("utf-8")
                    clean_text = _extract_text_from_xml(xml_content)
                    matches = re.findall(pattern, clean_text)
                    placeholders.update(matches)

            # Parser footer*.xml
            for name in docx_zip.namelist():
                if name.startswith("word/footer") and name.endswith(".xml"):
                    xml_content = docx_zip.read(name).decode("utf-8")
                    clean_text = _extract_text_from_xml(xml_content)
                    matches = re.findall(pattern, clean_text)
                    placeholders.update(matches)
    except Exception:
        pass

    return placeholders


def _extract_text_from_xml(xml_content: str) -> str:
    """
    Extrait le texte pur depuis le XML en supprimant les balises.

    Args:
        xml_content: Contenu XML brut

    Returns:
        Texte nettoyé sans balises XML
    """
    # Extraire uniquement le contenu des balises <w:t>
    text_pattern = r"<w:t[^>]*>([^<]*)</w:t>"
    text_parts = re.findall(text_pattern, xml_content)
    return "".join(text_parts)


def _replace_in_textboxes_xml(
    docx_path: Path, replacements: dict[str, str]
) -> None:
    """
    Remplace les placeholders dans les text boxes via XML direct.

    Gère la fragmentation des placeholders entre plusieurs balises <w:t>.

    Args:
        docx_path: Chemin vers le DOCX (sera modifié in-place)
        replacements: Dictionnaire {placeholder: valeur}
    """
    import shutil
    import tempfile

    try:
        temp_dir = Path(tempfile.mkdtemp())
        temp_docx = temp_dir / "temp.docx"

        # Extraire le ZIP
        with zipfile.ZipFile(docx_path, "r") as zip_read:
            zip_read.extractall(temp_dir)

        # Fichiers XML à modifier
        xml_files = ["word/document.xml"]
        xml_files.extend(
            [
                f"word/{name.name}"
                for name in (temp_dir / "word").iterdir()
                if name.name.startswith(("header", "footer"))
                and name.suffix == ".xml"
            ]
        )

        # Remplacer dans chaque fichier XML
        for xml_file in xml_files:
            xml_path = temp_dir / xml_file
            if xml_path.exists():
                content = xml_path.read_text(encoding="utf-8")
                
                # Méthode 1: Remplacement simple (pour placeholders non fragmentés)
                for field, value in replacements.items():
                    placeholder = f"{{{{{field}}}}}"
                    content = content.replace(placeholder, value)
                
                # Méthode 2: Gérer la fragmentation en nettoyant entre <w:t>
                content = _replace_fragmented_placeholders(content, replacements)
                
                xml_path.write_text(content, encoding="utf-8")

        # Recréer le ZIP
        with zipfile.ZipFile(temp_docx, "w", zipfile.ZIP_DEFLATED) as zip_write:
            for file in temp_dir.rglob("*"):
                if file.is_file() and file != temp_docx:
                    arcname = file.relative_to(temp_dir)
                    zip_write.write(file, arcname)

        # Remplacer le fichier original
        shutil.move(str(temp_docx), str(docx_path))
        shutil.rmtree(temp_dir, ignore_errors=True)

    except Exception:
        pass


def _replace_fragmented_placeholders(
    xml_content: str, replacements: dict[str, str]
) -> str:
    """
    Remplace les placeholders fragmentés dans le XML.

    Stratégie simple: parcourir les paragraphes, extraire le texte complet,
    faire les remplacements, et reconsolider dans une balise <w:t>.

    Args:
        xml_content: Contenu XML
        replacements: Dictionnaire de remplacements

    Returns:
        XML avec placeholders remplacés
    """
    # Pattern pour capturer un paragraphe complet
    para_pattern = r"(<w:p\b[^>]*>)(.*?)(</w:p>)"
    
    def process_paragraph(match):
        """Traite un paragraphe Word complet."""
        opening_tag = match.group(1)
        para_content = match.group(2)
        closing_tag = match.group(3)
        
        # Extraire tout le texte des balises <w:t>
        text_pattern = r"<w:t[^>]*>([^<]*)</w:t>"
        texts = re.findall(text_pattern, para_content)
        full_text = "".join(texts)
        
        # Vérifier si des placeholders sont présents
        has_placeholder = False
        new_text = full_text
        for field, value in replacements.items():
            placeholder = f"{{{{{field}}}}}"
            if placeholder in new_text:
                new_text = new_text.replace(placeholder, value)
                has_placeholder = True
        
        # Si des remplacements ont été effectués, reconsolider le texte
        if has_placeholder and new_text != full_text:
            # Garder la structure jusqu'au premier <w:r>, puis mettre le texte
            first_run = re.search(r"(<w:r\b[^>]*>.*?)(<w:t[^>]*>)", para_content)
            if first_run:
                # Tout avant le premier <w:t>
                before = para_content[:first_run.end()]
                # Trouver le </w:r> final
                last_run_end = para_content.rfind("</w:r>")
                after = para_content[last_run_end:] if last_run_end != -1 else ""
                
                # Reconstruire avec le texte consolidé
                new_para = f"{before}{new_text}</w:t>{after}"
                return f"{opening_tag}{new_para}{closing_tag}"
        
        return match.group(0)
    
    # Traiter tous les paragraphes
    result = re.sub(para_pattern, process_paragraph, xml_content, flags=re.DOTALL)
    
    return result


def _get_paragraph_text(paragraph) -> str:
    """
    Récupère le texte complet d'un paragraphe en gérant la fragmentation.

    Args:
        paragraph: Objet paragraphe python-docx

    Returns:
        Texte complet reconstruit
    """
    return "".join(run.text for run in paragraph.runs)


def _replace_in_table(table, replacements: dict[str, str]) -> None:
    """
    Remplace les placeholders dans un tableau.

    Args:
        table: Objet table python-docx
        replacements: Dictionnaire {placeholder: valeur}
    """
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _replace_in_paragraph(paragraph, replacements)
            # Gérer les tables imbriquées
            for nested_table in cell.tables:
                _replace_in_table(nested_table, replacements)


def _replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    """
    Remplace les placeholders dans un paragraphe en gérant fragmentation.

    Args:
        paragraph: Objet paragraphe python-docx
        replacements: Dictionnaire {placeholder: valeur}
    """
    from app.core.config import FIELD_PATTERN

    # Reconstituer le texte complet
    full_text = _get_paragraph_text(paragraph)
    matches = re.findall(FIELD_PATTERN, full_text)

    if not matches:
        return

    # Remplacer les placeholders
    new_text = full_text
    for match in matches:
        if match in replacements:
            placeholder = f"{{{{{match}}}}}"
            new_text = new_text.replace(placeholder, replacements[match])

    # Appliquer le nouveau texte si modifié
    if new_text != full_text:
        # Garder le formatage du premier run, supprimer les autres
        if paragraph.runs:
            paragraph.runs[0].text = new_text
            for i in range(len(paragraph.runs) - 1, 0, -1):
                paragraph.runs[i].text = ""


def insert_image_in_docx(
    docx_path: Path,
    image_path: Path,
    placeholder: str,
    output_path: Path,
    width_inches: float = 2.0,
) -> None:
    """
    Remplace un placeholder par une image dans un DOCX.

    Args:
        docx_path: Chemin vers le template DOCX
        image_path: Chemin vers l'image à insérer
        placeholder: Nom du placeholder (ex: 'img_logo')
        output_path: Chemin de sortie
        width_inches: Largeur de l'image en inches
    """
    doc = Document(docx_path)
    placeholder_str = f"{{{{{placeholder}}}}}"

    for paragraph in doc.paragraphs:
        full_text = _get_paragraph_text(paragraph)
        if placeholder_str in full_text:
            # Supprimer tout le texte des runs
            for run in paragraph.runs:
                run.text = ""
            # Ajouter l'image
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(width_inches))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    full_text = _get_paragraph_text(paragraph)
                    if placeholder_str in full_text:
                        # Supprimer tout le texte des runs
                        for run in paragraph.runs:
                            run.text = ""
                        # Ajouter l'image
                        run = paragraph.add_run()
                        run.add_picture(
                            str(image_path), width=Inches(width_inches)
                        )

    doc.save(output_path)

